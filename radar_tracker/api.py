"""Rotas HTTP: interface web, API REST do painel e endpoint do OwnTracks."""

import hmac
import io
import logging
import time
from datetime import datetime, timedelta

from flask import Blueprint, abort, current_app, jsonify, render_template, request, send_file
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from . import roteamento
from .frota import MotoboyNaoEncontrado

log = logging.getLogger(__name__)
bp = Blueprint("radar", __name__)


def _frota():
    return current_app.extensions["radar"]["frota"]


def _banco():
    return current_app.extensions["radar"]["banco"]


def _cfg():
    return current_app.extensions["radar"]["config"]


def _json():
    dados = request.get_json(silent=True)
    if not isinstance(dados, dict):
        abort(400, description="Corpo JSON inválido.")
    return dados


def _campo(dados, nome, tipo=str):
    try:
        valor = tipo(dados[nome])
    except (KeyError, TypeError, ValueError):
        abort(400, description=f"Campo '{nome}' ausente ou inválido.")
    if tipo is str:
        valor = valor.strip()
        if not valor:
            abort(400, description=f"Campo '{nome}' não pode ser vazio.")
    return valor


@bp.errorhandler(MotoboyNaoEncontrado)
def _motoboy_inexistente(_erro):
    return jsonify({"ok": 0, "erro": "Motoboy não encontrado."}), 404


@bp.app_errorhandler(400)
def _requisicao_invalida(erro):
    return jsonify({"ok": 0, "erro": erro.description}), 400


# --- Interface ---
@bp.route("/")
def index():
    cfg = _cfg()
    return render_template("index.html", empresa=cfg.NOME_EMPRESA, base_lat=cfg.BASE_LAT, base_lon=cfg.BASE_LON)


@bp.route("/api/health")
def health():
    return jsonify({"status": "ok"})


@bp.route("/api/motoboys")
def motoboys():
    return jsonify([{"id": m["id"], "nome": m["nome"], "cor": m["cor"]} for m in _frota().motoboys.values()])


# --- Frota / rota ---
@bp.route("/api/estado")
def estado():
    return jsonify(_frota().snapshot())


@bp.route("/api/add", methods=["POST"])
def adicionar_parada():
    d = _json()
    ok = _frota().adicionar_parada(_campo(d, "id", int), _campo(d, "nome"), _campo(d, "lat", float), _campo(d, "lon", float))
    return jsonify({"ok": int(ok)})


@bp.route("/api/rm", methods=["POST"])
def remover_parada():
    d = _json()
    ok = _frota().remover_parada(_campo(d, "id", int), _campo(d, "idx", int))
    return jsonify({"ok": int(ok)})


@bp.route("/api/otimizar", methods=["POST"])
def otimizar():
    ok = _frota().otimizar(_campo(_json(), "id", int))
    return jsonify({"ok": int(ok)})


@bp.route("/api/start", methods=["POST"])
def iniciar():
    d = _json()
    ok = _frota().iniciar(_campo(d, "id", int), str(d.get("nota", "")).strip())
    return jsonify({"ok": int(ok)})


@bp.route("/api/end", methods=["POST"])
def finalizar():
    viagem = _frota().finalizar(_campo(_json(), "id", int))
    if viagem:
        _banco().registrar_viagem(**viagem)
    return jsonify({"ok": int(bool(viagem))})


@bp.route("/api/calc", methods=["POST"])
def calcular():
    paradas = _json().get("stops", [])
    try:
        paradas = [{"lat": float(p["lat"]), "lon": float(p["lon"])} for p in paradas]
    except (KeyError, TypeError, ValueError):
        abort(400, description="Lista de paradas inválida.")
    cfg = _cfg()
    return jsonify(roteamento.calcular_rota(paradas, {"lat": cfg.BASE_LAT, "lon": cfg.BASE_LON}, cfg))


@bp.route("/api/buscar")
def buscar():
    q = request.args.get("q", "").strip()
    return jsonify(roteamento.buscar_endereco(q, _cfg()) if q else [])


# --- GPS (OwnTracks, modo HTTP) ---
def _owntracks_autorizado():
    """Se OWNTRACKS_TOKEN estiver definido, exige-o como senha (HTTP Basic) ou header X-Token."""
    token = _cfg().OWNTRACKS_TOKEN
    if not token:
        return True
    auth = request.authorization
    enviado = (auth.password if auth else None) or request.headers.get("X-Token", "")
    return hmac.compare_digest(str(enviado), token)


def _numero(valor):
    try:
        return float(valor)
    except (TypeError, ValueError):
        return None


@bp.route("/api/gps/owntracks", methods=["POST"])
def gps_owntracks():
    if not _owntracks_autorizado():
        return jsonify({"erro": "Não autorizado."}), 401
    d = request.get_json(silent=True) or {}
    if d.get("_type") != "location":
        return jsonify([])  # OwnTracks espera uma lista (vazia) como resposta
    frota = _frota()
    mid = frota.id_por_tid(d.get("tid", ""))
    lat, lon = _numero(d.get("lat")), _numero(d.get("lon"))
    if mid is None or lat is None or lon is None:
        log.info("OwnTracks: leitura ignorada (tid=%r).", d.get("tid"))
        return jsonify([])
    ts = int(d.get("tst") or time.time())
    vel = frota.atualizar_posicao(mid, lat, lon, ts, precisao_m=_numero(d.get("acc")), vel_dispositivo=_numero(d.get("vel")))
    if vel is None:
        log.info("OwnTracks: leitura descartada pelo filtro (moto %s, acc=%s).", mid, d.get("acc"))
    else:
        _banco().registrar_posicao(mid, lat, lon, vel, ts)
    return jsonify([])


@bp.route("/api/tracklog")
def tracklog():
    mid = request.args.get("id", 0, type=int)
    data = request.args.get("data")
    if data:
        try:
            inicio = datetime.strptime(data, "%Y-%m-%d")
        except ValueError:
            abort(400, description="Data inválida (use AAAA-MM-DD).")
        ts_inicio, ts_fim = int(inicio.timestamp()), int((inicio + timedelta(days=1)).timestamp())
    else:
        ts_fim = int(time.time()) + 1
        ts_inicio = ts_fim - 86400
    return jsonify(_banco().trilha(mid, ts_inicio, ts_fim))


# --- Clientes ---
@bp.route("/api/clientes")
def clientes():
    return jsonify(_banco().listar_clientes())


@bp.route("/api/save_cli", methods=["POST"])
def salvar_cliente():
    d = _json()
    cid = _banco().salvar_cliente(_campo(d, "nome"), _campo(d, "lat", float), _campo(d, "lon", float))
    return jsonify({"ok": 1, "id": cid})


@bp.route("/api/update_client", methods=["POST"])
def atualizar_cliente():
    d = _json()
    _banco().atualizar_cliente(_campo(d, "id", int), _campo(d, "nome"), str(d.get("telefone", "")).strip())
    return jsonify({"ok": 1})


@bp.route("/api/update_color", methods=["POST"])
def atualizar_cor():
    d = _json()
    cor = _campo(d, "cor")
    if not (cor.startswith("#") and len(cor) == 7):
        abort(400, description="Cor deve estar no formato #rrggbb.")
    _banco().atualizar_cor_cliente(_campo(d, "id", int), cor)
    return jsonify({"ok": 1})


@bp.route("/api/apagar_cliente", methods=["POST"])
def apagar_cliente():
    _banco().apagar_cliente(_campo(_json(), "id", int))
    return jsonify({"ok": 1})


# --- Histórico e relatórios ---
@bp.route("/api/get_historico")
def historico():
    return jsonify(_banco().historico(request.args.get("inicio"), request.args.get("fim")))


@bp.route("/baixar_excel")
def baixar_excel():
    linhas = _banco().historico(request.args.get("inicio"), request.args.get("fim"))
    colunas = [
        ("data_viagem", "Data"),
        ("motoboy", "Motoboy"),
        ("numero_nota", "Nota / Pedido"),
        ("destinos", "Destinos"),
        ("hora_saida", "Saída"),
        ("hora_chegada", "Chegada"),
    ]
    wb = Workbook()
    ws = wb.active
    ws.title = "Histórico"
    ws.append([titulo for _, titulo in colunas])
    for celula in ws[1]:
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = PatternFill("solid", fgColor="2C3E50")
    for linha in linhas:
        ws.append([linha.get(chave) or "" for chave, _ in colunas])
    for letra, largura in zip("ABCDEF", (12, 20, 16, 60, 10, 10)):
        ws.column_dimensions[letra].width = largura
    ws.freeze_panes = "A2"

    saida = io.BytesIO()
    wb.save(saida)
    saida.seek(0)
    nome = f"historico_entregas_{datetime.now():%Y%m%d}.xlsx"
    return send_file(
        saida,
        as_attachment=True,
        download_name=nome,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
